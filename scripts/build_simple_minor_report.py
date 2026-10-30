#!/usr/bin/env python3
"""Generate simple_minor_report.docx from Minor-report-ECE template structure."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "simple_minor_report.docx"

PROJECT_TITLE = (
    "Adaptive CUDA Graph Admission for Variable-Shape Text Encoder Inference"
)


def set_run_font(run, size_pt: int, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.bold = bold


def add_heading(doc: Document, text: str, level: int) -> None:
    """level 1 -> 16pt, level 2 -> 14pt, level 3 -> 12pt bold."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, 16, bold=True)
    elif level == 2:
        set_run_font(run, 14, bold=True)
    else:
        set_run_font(run, 12, bold=True)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, 12)


def main() -> None:
    doc = Document()

    # --- Title page ---
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(PROJECT_TITLE)
    set_run_font(r, 18, bold=True)

    lines = [
        "Minor Specialization Project Report",
        "Course Code: ECE_4191",
        "",
        "Submitted in Partial Fulfilment of the Requirements for the Award of the Degree of",
        "Bachelor of Technology",
        "In",
        "Electronics and Communication Engineering",
        "",
        "Submitted by",
        "Name of the Student _______________________________________________________",
        "Registration Number _____________________________________________________",
        "",
        "School of Electrical Engineering",
        "Manipal Institute of Technology Bengaluru",
        "July-December 2026",
    ]
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, 12, bold=("Minor Specialization" in line or line == PROJECT_TITLE))

    doc.add_page_break()

    # --- Certificate ---
    add_heading(doc, "Certificate", 1)
    add_body(
        doc,
        f'This is to certify that Mr./Ms. ______________________________, bearing Registration '
        f'Number ________________, has successfully completed the Minor project titled '
        f'"{PROJECT_TITLE}" under my supervision during the academic year 2026 as a part of '
        f'the requirements for the Minor Specialization in partial fulfilment of the requirements '
        f'for the award of the B.Tech. Degree.',
    )
    add_body(
        doc,
        "The student has carried out the project work satisfactorily and submitted the project "
        "report as per the prescribed guidelines. The work presented is original and has been "
        "completed to the best of the student's knowledge and effort.",
    )
    doc.add_paragraph()
    add_body(doc, "Name")
    add_body(doc, "Project Mentor")
    add_body(doc, "SoEE, MIT")
    doc.add_page_break()

    # --- Contents ---
    add_heading(doc, "Contents", 1)
    toc = [
        ("01", "Introduction", "1"),
        ("02", "Literature survey", "3"),
        ("03", "Workflow / methodology", "8"),
        ("04", "Results", "14"),
        ("05", "Conclusion and future scope", "18"),
        ("", "Reference", "19"),
    ]
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    for i, label in enumerate(["Sl.No.", "Contents", "Page No."]):
        run = hdr[i].paragraphs[0].add_run(label)
        set_run_font(run, 12, bold=True)
    for sl, title, pg in toc:
        row = table.add_row().cells
        for i, val in enumerate([sl, title, pg]):
            run = row[i].paragraphs[0].add_run(val)
            set_run_font(run, 12)
    doc.add_page_break()

    # --- Chapter 1 ---
    add_heading(doc, "Chapter 1: Introduction", 1)

    add_heading(doc, "1.1 Background", 2)
    add_body(
        doc,
        "Deep learning inference on NVIDIA GPUs is dominated by repeated neural network forward "
        "passes. PyTorch executes each forward eagerly: Python dispatches many small CUDA kernel "
        "launches (matrix multiplies, attention, normalization). For encoder workloads with "
        "moderate compute per token, launch and framework overhead can be significant. CUDA "
        "graphs record a fixed sequence of GPU operations once and replay them with reduced launch "
        "cost. Graphs require static operator topology and tensor shapes; inputs are updated via "
        "in-place copies into captured buffers.",
    )

    add_heading(doc, "1.2 Problem Statement", 2)
    add_body(
        doc,
        "Real text encoders see variable batch sizes and sequence lengths. Each distinct "
        "(batch, sequence length) pair needs a separate captured graph. Naively capturing every "
        "new shape incurs warmup and capture latency, consumes GPU memory, and interacts badly "
        "with bounded graph caches: under PyTorch shared graph pool semantics, evicting one graph "
        "can invalidate the entire cache. The problem is to decide online which shapes should be "
        "admitted to a limited graph cache so that end-to-end latency improves versus always "
        "running eager forwards.",
    )

    add_heading(doc, "1.3 Motivation", 2)
    add_body(
        doc,
        "Production systems (embedding APIs, rerankers, small language model encoders) often "
        "multiplex heterogeneous requests. Prior work on CUDA graphs in LLM serving focuses on "
        "decode-phase piecewise capture. Less attention is given to admission control for "
        "variable-shape encoder grids under explicit memory budgets. A reproducible research "
        "prototype helps quantify when graphs help and when capture storms hurt.",
    )

    add_heading(doc, "1.4 Objectives", 2)
    add_bullets(
        doc,
        [
            "Implement correct eager and per-shape CUDA graph capture/replay for Hugging Face "
            "encoder models.",
            "Calibrate per-shape eager, replay, and capture costs on the target GPU.",
            "Simulate seeded workloads (recurring, many rare shapes, shifting hot sets) with "
            "four routing policies.",
            "Compare total workload time, capture counts, memory, and correctness against an "
            "always-eager baseline.",
            "Provide an online adaptive admission policy using causal demand estimates and "
            "calibrated economics.",
        ],
    )

    add_heading(doc, "1.5 Scope of Work", 2)
    add_body(
        doc,
        "The scope is a research prototype in Python/PyTorch, not integration into a production "
        "serving stack. Models are loaded via Hugging Face AutoModel; forwards use last_hidden_state "
        "with use_cache=False. Experiments use a discrete shape grid from YAML configuration. "
        "Bundled configs cover MiniLM-L6-v2 (default), Qwen2.5-1.5B, and SmolLM2-1.7B. Optional "
        "padding-to-graph routing is implemented but disabled in the main baseline.",
    )

    add_heading(doc, "1.6 Organization of the Report", 2)
    add_body(
        doc,
        "Chapter 2 reviews related CUDA graph and serving work. Chapter 3 describes system "
        "architecture, policies, and experimental methodology. Chapter 4 presents setup, results, "
        "and discussion on an RTX 5060 Laptop GPU. Chapter 5 concludes and outlines future work.",
    )

    doc.add_page_break()

    # --- Chapter 2 ---
    add_heading(doc, "Chapter 2: Literature Review", 1)

    add_heading(doc, "2.1 Existing Systems", 2)
    add_body(
        doc,
        "PyTorch exposes torch.cuda.CUDAGraph and graph pool handles for capture and replay. "
        "torch.compile can fuse and graph regions automatically. TensorRT and similar compilers "
        "also exploit graph-style execution. LLM inference frameworks such as vLLM employ "
        "CUDA graphs for decode and partial capture strategies tuned to autoregressive serving.",
    )

    add_heading(doc, "2.2 Review of Related Work", 2)
    add_body(
        doc,
        "CUDA graph literature emphasizes capture mechanics, kernel fusion, and static-shape "
        "benchmarks. Serving papers prioritize throughput under continuous batching. Cache "
        "replacement for compiled artifacts appears in other domains but is less standardized "
        "for multi-shape encoder graphs with shared pool invalidation on partial eviction.",
    )

    add_heading(doc, "2.3 Research Gap", 2)
    add_body(
        doc,
        "There is a gap between always capturing new shapes and offline profiling with fixed "
        "precapture sets: online admission under a bounded cache and GPU memory budget, using "
        "only causal history and device-calibrated capture versus replay savings, for "
        "variable-shape encoder inference.",
    )

    add_heading(doc, "2.4 Summary", 2)
    add_body(
        doc,
        "Existing tools provide graph mechanisms; this project contributes policy-level admission "
        "and comparative evaluation on reproducible traces rather than new kernel fusion.",
    )

    doc.add_page_break()

    # --- Chapter 3 ---
    add_heading(doc, "Chapter 3: Methodology / System Design", 1)

    add_heading(doc, "3.1 Proposed System", 2)
    add_body(
        doc,
        "The system is a trace-driven simulator. Each request specifies (batch, seq_len). A "
        "policy decides eager execution, graph replay, or capture. A GraphCache stores up to "
        "max_graph_cache shapes subject to mem_budget_mb. Device calibration JSON supplies "
        "p50 eager, replay, and capture latencies plus memory delta per shape.",
    )

    add_heading(doc, "3.2 System Architecture", 2)
    add_bullets(
        doc,
        [
            "EncoderModel: AutoTokenizer + AutoModel on CUDA.",
            "CapturedShapeGraph: static buffers, warmup, capture, replay with input copy and output clone.",
            "GraphCache: LRU metadata with full-pool flush on eviction.",
            "Policies: always_eager, lru_on_first_use, calibration_fixed, adaptive_admission.",
            "Runner: executes 200-request traces, logs JSONL with per-request latency and reasons.",
            "analyze_results.py: aggregates total_s, p50, p95, captures, correctness.",
        ],
    )

    add_heading(doc, "3.3 Block Diagram", 2)
    add_body(
        doc,
        "Trace request → Policy.decide → [capture if needed] → eager forward OR graph replay → "
        "optional correctness check vs eager → append JSONL record. Calibration runs offline "
        "once per (device, model, shape grid) before policy experiments.",
    )

    add_heading(doc, "3.4 Hardware Design", 2)
    add_body(
        doc,
        "No custom ASIC or FPGA design is involved. The experimental platform is a commodity "
        "NVIDIA GPU (development and results on NVIDIA GeForce RTX 5060 Laptop GPU, 8 GB class). "
        "Software requires CUDA-capable drivers (including WSL2 when applicable). Memory budget "
        "parameters in software mirror VRAM constraints for graph storage.",
    )

    add_heading(doc, "3.5 Software Design", 2)
    add_body(
        doc,
        "Repository layout: src/acg/ (model, graph_pool, policies, runner, calibration, traces); "
        "scripts/ (verify_one_shape, run_experiment, analyze_results, run_large_model_benchmark); "
        "configs/*.yaml; calibration/*.json; traces/ and results/. Configuration selects model_id, "
        "dtype, batch_sizes, seq_lengths, policy hyperparameters, and output subdirectories.",
    )

    add_heading(doc, "3.6 Flowchart", 2)
    add_body(
        doc,
        "For adaptive_admission on cache miss: estimate P(shape) from the last W requests; "
        "compute expected_savings = P × horizon × (eager_p50 − replay_p50); capture if "
        "expected_savings > capture_p50 × margin and memory budget allows; else run eager.",
    )

    add_heading(doc, "3.7 Working Principle", 2)
    add_body(
        doc,
        "Policy 1 (always_eager): baseline. Policy 2 (lru_on_first_use): capture on every miss. "
        "Policy 3 (calibration_fixed): observe first K requests eagerly, precapture top M shapes, "
        "no further captures. Policy 4 (adaptive_admission): online admission test using rolling "
        "window (default 50), horizon (default 80), margin (default 1.25). Workloads: recurring "
        "(80% hot shapes), many_rare (uniform over grid), shift (hot set changes mid-trace).",
    )

    doc.add_page_break()

    # --- Chapter 4 ---
    add_heading(doc, "Chapter 4: Implementation and Results", 1)

    add_heading(doc, "4.1 Experimental Setup", 2)
    add_body(
        doc,
        "Platform: Linux WSL2, PyTorch 2.14+cu130, RTX 5060 Laptop GPU. Default model: "
        "sentence-transformers/all-MiniLM-L6-v2, float32, shape grid batch {1,4,8} × seq "
        "{32,64,128,256}, max_graph_cache=4, mem_budget_mb=768, 200 requests per trace, "
        "3 experiment repeats. Calibration: rtx5060.json. Large-model config: Qwen2.5-1.5B bf16, "
        "smaller grid, separate calibration and results under qwen2_5_1_5b/.",
    )

    add_heading(doc, "4.2 Testing Procedure", 2)
    add_bullets(
        doc,
        [
            "verify_one_shape.py: single-shape eager vs graph output match.",
            "run_experiment.py --calibrate-only: write calibration JSON.",
            "run_experiment.py --generate-traces: traces + full policy matrix.",
            "analyze_results.py: tabulate mean±stdev of total_workload_s and other metrics.",
            "Correctness: sampled every 20 requests (outputs_match vs eager).",
        ],
    )

    add_heading(doc, "4.3 Results and Analysis", 2)
    add_body(
        doc,
        "Primary metric: total_workload_s (seconds for 200 requests, including capture and warmup). "
        "MiniLM aggregated results (three repeats, RTX 5060):",
    )
    add_bullets(
        doc,
        [
            "recurring — always_eager: 0.859±0.36 s; calibration_fixed: 0.725±0.01 s (~15.7% faster); "
            "adaptive_admission: 0.934±0.50 s; lru_on_first_use: 2.632±0.10 s (much slower).",
            "many_rare — always_eager: 0.937±0.04 s; adaptive_admission: 0.855±0.02 s (~8.8% faster); "
            "calibration_fixed: 1.217 s; lru: 4.783 s.",
            "shift — always_eager: 1.216 s; adaptive_admission: 1.015 s (~16.5% faster); "
            "calibration_fixed: 1.427 s; lru: 3.110 s.",
            "Capture counts (mean): LRU 59–147; calibration_fixed 3; adaptive 3–7 vs 0 for eager.",
            "Correctness rate: 1.0 on completed MiniLM runs.",
        ],
    )
    add_body(
        doc,
        "Per-shape calibration shows replay p50 below eager p50 by roughly 0.3–0.9 ms per hit on "
        "small MiniLM shapes; savings compound when shape probability is high. Qwen2.5-1.5B "
        "partial runs (one repeat): calibration_fixed improved recurring (~26%) and many_rare (~5%) "
        "vs eager; adaptive and LRU runs in stored artifacts were slower—full repeats with "
        "model-specific calibration are recommended before final presentation.",
    )

    add_heading(doc, "4.4 Performance Evaluation", 2)
    add_body(
        doc,
        "Performance is evaluated against always_eager, not against impractical unlimited capture. "
        "LRU-on-first-use demonstrates negative results: capture storms dominate when shapes are "
        "rare or cache flushes are frequent. Adaptive admission reduces captures while beating "
        "eager on many_rare and shift. Fixed precapture wins stable recurring workloads but degrades "
        "when the hot set shifts after the observation window.",
    )

    add_heading(doc, "4.5 Discussion", 2)
    add_body(
        doc,
        "The prototype is model-agnostic via YAML (any Hugging Face AutoModel encoder forward) but "
        "requires per-device, per-model calibration and a shape grid that fits VRAM. Graph replay "
        "includes copy and clone costs in measurements. Shared graph pool flush on eviction is "
        "faithful to PyTorch semantics and explains part of LRU pathology. The novel contribution "
        "is the admission rule and empirical policy comparison on reproducible traces, not raw "
        "graph speedups alone.",
    )

    doc.add_page_break()

    # --- Chapter 5 ---
    add_heading(doc, "Chapter 5: Conclusion and Future Scope", 1)

    add_heading(doc, "5.1 Conclusion", 2)
    add_body(
        doc,
        "This minor project implemented and evaluated adaptive CUDA graph admission for "
        "variable-shape text encoder inference. CUDA graphs reduce per-forward overhead when "
        "shapes repeat, but uncritical capture hurts total latency. Device-calibrated adaptive "
        "admission and fixed precapture both outperform naive LRU and always-eager on selected "
        "traces on RTX 5060 with MiniLM; the best policy depends on workload stationarity.",
    )

    add_heading(doc, "5.2 Achievements", 2)
    add_bullets(
        doc,
        [
            "End-to-end reproducible pipeline: calibration, traces, four policies, JSONL logs, analysis.",
            "Demonstrated up to ~16.5% total time reduction vs eager (shift trace, adaptive).",
            "Documented cache pool eviction behavior and correctness sampling.",
            "Extended configs for ~1–2B models (Qwen, SmolLM2) with separate artifact paths.",
        ],
    )

    add_heading(doc, "5.3 Future Enhancements", 2)
    add_bullets(
        doc,
        [
            "Complete Qwen/SmolLM benchmarks with multiple repeats and tuned admission hyperparameters.",
            "Enable and ablate padding_enabled routing to captured larger shapes.",
            "Integrate with continuous batching or torch.compile for hybrid strategies.",
            "Learned or skew-aware admission without full grid calibration.",
            "Deploy prototype behind a small REST embedding service for live trace collection.",
        ],
    )

    doc.add_page_break()

    add_heading(doc, "Reference", 1)
    refs = [
        "PyTorch Documentation, CUDA Graphs, pytorch.org, accessed 2026.",
        "NVIDIA CUDA C Programming Guide, Graph Management, docs.nvidia.com, accessed 2026.",
        "Hugging Face Transformers, AutoModel and tokenizer APIs, huggingface.co, accessed 2026.",
        "Kwon et al., Efficient Memory Management for LLM Serving with PagedAttention (vLLM), 2023.",
        "Sentence Transformers, all-MiniLM-L6-v2 model card, huggingface.co, accessed 2026.",
        "Qwen Team, Qwen2.5 model family technical report, 2024.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"[{i}] {ref}")
        set_run_font(run, 12)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
